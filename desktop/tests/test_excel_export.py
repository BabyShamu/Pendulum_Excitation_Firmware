"""Tests for normalized experiment workbook output."""
import hashlib
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from pendulum_workbench.experiments.excel_export import (
    EXPERIMENT_COLUMNS, export_experiment_workbook
)


class ExperimentExcelExportTests(unittest.TestCase):
    def test_three_sheets_numeric_columns_and_simulation_status(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            folder = root / "2026-10-24_FS"
            folder.mkdir()
            manifest = {
                "experiment_id": "id-123", "name": "FS_test", "created_at": "2026-10-24T10:00:00",
                "completion_status": "COMPLETED", "peak_absolute_angle_deg": 10.2,
                "experiment": {
                    "experiment_mode": "FREE_SWING",
                    "measured_initial_angle_deg": 9.84,
                    "requested_duration_s": 20,
                    "actual_elapsed_duration_s": 20.1,
                    "commanded_amplitude_mm": None,
                },
            }
            (folder / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
            saved = export_experiment_workbook([(folder, manifest)], root)
            workbook = load_workbook(saved, data_only=True)
            self.assertEqual(workbook.sheetnames, ["Experiments", "Simulations", "Comparisons"])
            sheet = workbook["Experiments"]
            headers = [cell.value for cell in sheet[1]]
            self.assertEqual(headers, EXPERIMENT_COLUMNS)
            row = {key: sheet.cell(2, i + 1).value for i, key in enumerate(headers)}
            self.assertEqual(row["Type"], "Free Swing")
            self.assertEqual(row["Initial Angle (deg)"], 9.84)
            self.assertIsInstance(row["Actual Duration (s)"], float)
            self.assertIsNone(row["Amplitude (mm)"])
            self.assertEqual(row["Simulation Status"], "Not simulated")
            self.assertEqual(row["Comparison Status"], "Not compared")

    def test_saved_comparison_metrics_are_numeric(self):
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            folder = root / "experiment"
            sim = folder / "simulation"
            sim.mkdir(parents=True)
            manifest = {"experiment_id": "id-456", "name": "sine",
                        "experiment": {"experiment_mode": "VERTICAL_SINUSOIDAL_EXCITATION",
                                       "commanded_amplitude_mm": 12,
                                       "commanded_frequency_hz": 1.8}}
            (folder / "manifest.json").write_text(json.dumps(manifest), encoding="utf-8")
            input_file = folder / "simulation_input.json"
            input_file.write_text('{"experiment_id":"id-456"}', encoding="utf-8")
            digest = hashlib.sha256(input_file.read_bytes()).hexdigest()
            (sim / "metadata.json").write_text(json.dumps({
                "schema_version": 1, "status": "complete",
                "experiment_id": "id-456", "fingerprint": {"input_sha256": digest},
                "simulation_model": "Pendulum_Simscape_Model"
            }), encoding="utf-8")
            for filename in ("simulation.mat", "simulation.csv", "pendulum_angle.png",
                             "pivot_position.png", "simulation_overview.fig"):
                (sim / filename).write_bytes(b"test")
            (sim / "comparison_summary.json").write_text(json.dumps({
                "status": "Preliminary", "angle_rmse_deg": 0.55, "pivot_rmse_mm": 0.2,
                "measured_peak_absolute_angle_deg": 30,
                "simulated_peak_absolute_angle_deg": 29
            }), encoding="utf-8")
            exported = export_experiment_workbook([(folder, manifest)], root)
            wb = load_workbook(exported, data_only=True)
            self.assertEqual(wb["Experiments"].cell(2, EXPERIMENT_COLUMNS.index("Angle RMSE (deg)") + 1).value, 0.55)
            self.assertEqual(wb["Experiments"].cell(2, EXPERIMENT_COLUMNS.index("Simulation Status") + 1).value, "Saved")
            self.assertEqual(wb["Comparisons"].cell(2, 3).value, "Preliminary")


if __name__ == "__main__":
    unittest.main()
