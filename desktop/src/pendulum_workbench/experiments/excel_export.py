"""Structured Excel index of experiments, Simscape runs and comparisons."""
from __future__ import annotations

import json
import math
import os
from pathlib import Path

from pendulum_workbench.experiments.simulation_cache import inspect_simulation

EXPERIMENT_COLUMNS = [
    "Experiment ID", "Date", "Name", "Type", "Completion", "Archived",
    "Initial Angle (deg)", "Initial Omega (rad/s)",
    "Amplitude (mm)", "Frequency (Hz)", "Feedback Phase (deg)",
    "Target Angle (deg)", "Requested Duration (s)", "Actual Duration (s)",
    "Measured Peak Angle (deg)", "Measured Peak Time (s)",
    "Time to Target (s)", "Samples", "Raw Lines",
    "Simulation Status", "Comparison Status", "Angle RMSE (deg)",
    "Pivot RMSE (mm)", "Notes", "Experiment Folder",
]
SIMULATION_COLUMNS = [
    "Experiment ID", "Name", "Simulation Status", "Model",
    "Simulation Created", "Input Fingerprint", "Model Fingerprint",
    "Dictionary Fingerprint", "Simulation Folder", "Simulation MAT",
]
COMPARISON_COLUMNS = [
    "Experiment ID", "Name", "Comparison Status", "Angle RMSE (deg)",
    "Pivot RMSE (mm)", "Measured Peak (deg)", "Simulated Peak (deg)",
    "Alignment Note", "Comparison Folder",
]


def _number(value):
    if isinstance(value, bool):
        return None
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def _read_json(path: Path):
    try:
        result = json.loads(path.read_text(encoding="utf-8"))
        return result if isinstance(result, dict) else {}
    except (OSError, ValueError):
        return {}


def _comparison(directory: Path):
    # The comparison summary is deliberately independent of cache validity.
    # A change to the model makes both simulation and comparison stale.
    return _read_json(directory / "simulation" / "comparison_summary.json")


def export_experiment_workbook(records, root: Path) -> Path:
    """Create/replace an Excel workbook without changing recorded experiments."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    root = Path(root)
    root.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    experiments = workbook.active
    experiments.title = "Experiments"
    simulations = workbook.create_sheet("Simulations")
    comparisons = workbook.create_sheet("Comparisons")
    experiments.append(EXPERIMENT_COLUMNS)
    simulations.append(SIMULATION_COLUMNS)
    comparisons.append(COMPARISON_COLUMNS)

    for directory, manifest in records:
        directory = Path(directory)
        details = manifest.get("experiment") or {}
        if not isinstance(details, dict):
            details = {}
        experiment_id = str(manifest.get("experiment_id") or "")
        name = str(manifest.get("name") or "")
        mode = {
            "FREE_SWING": "Free Swing",
            "VERTICAL_SINUSOIDAL_EXCITATION": "Sinusoidal",
            "AUTO_RESONANCE_PARAMETRIC": "Auto Resonance",
            "PHASE_TRACKED_SINE_AUTO_RESONANCE": "Phase Tracked",
        }.get(details.get("experiment_mode"), str(details.get("experiment_mode") or "Recording"))
        status = inspect_simulation(directory)
        metadata = status.metadata
        simulation = directory / "simulation"
        comparison = _comparison(directory)
        comparison_status = (
            comparison.get("status", "Not compared") if status.state == "ready"
            else "Stale" if comparison else "Not compared"
        )
        experiments.append([
            experiment_id,
            str(manifest.get("created_at") or ""),
            name, mode, str(manifest.get("completion_status") or ""),
            "Yes" if manifest.get("archived") else "No",
            _number(details.get("measured_initial_angle_deg")),
            _number(details.get("initial_angular_velocity_rad_s")),
            _number(details.get("commanded_amplitude_mm")),
            _number(details.get("commanded_frequency_hz")),
            _number(details.get("commanded_phase_deg")),
            _number(details.get("target_angle_deg")),
            _number(details.get("requested_duration_s") or details.get("commanded_duration_s")),
            _number(details.get("actual_elapsed_duration_s")),
            _number(manifest.get("peak_absolute_angle_deg")),
            _number(manifest.get("peak_angle_elapsed_s")),
            _number(details.get("time_to_target_s")),
            _number(manifest.get("sample_count")),
            _number(manifest.get("raw_line_count")),
            {"ready": "Saved", "missing": "Not simulated", "stale": "Stale", "incomplete": "Incomplete"}[status.state],
            comparison_status,
            _number(comparison.get("angle_rmse_deg")),
            _number(comparison.get("pivot_rmse_mm")),
            str(manifest.get("notes") or details.get("notes") or ""),
            str(directory),
        ])
        fingerprint = metadata.get("fingerprint") or {}
        simulations.append([
            experiment_id, name,
            {"ready": "Saved", "missing": "Not simulated", "stale": "Stale", "incomplete": "Incomplete"}[status.state],
            metadata.get("simulation_model"),
            metadata.get("created_at"),
            fingerprint.get("input_sha256"),
            fingerprint.get("model_sha256"),
            fingerprint.get("dictionary_sha256"),
            str(simulation),
            str(simulation / "simulation.mat") if (simulation / "simulation.mat").is_file() else "",
        ])
        comparisons.append([
            experiment_id, name, comparison_status,
            _number(comparison.get("angle_rmse_deg")),
            _number(comparison.get("pivot_rmse_mm")),
            _number(comparison.get("measured_peak_absolute_angle_deg")),
            _number(comparison.get("simulated_peak_absolute_angle_deg")),
            comparison.get("alignment_note"),
            str(simulation),
        ])

    for sheet in workbook:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="DAEAE2")
            cell.font = Font(bold=True, color="254D3C")
            cell.alignment = Alignment(wrap_text=True)
        sheet.row_dimensions[1].height = 36
        for index, column in enumerate(sheet.columns, start=1):
            letter = get_column_letter(index)
            # Keep long paths and fingerprints bounded and readable.
            max_length = max(len(str(cell.value or "")) for cell in column)
            sheet.column_dimensions[letter].width = min(44, max(15, max_length + 2))
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                if cell.data_type == "n":
                    cell.number_format = "0.0000"

    destination = root / "Pendulum_Experiments.xlsx"
    temp = root / ".Pendulum_Experiments.tmp.xlsx"
    try:
        workbook.save(temp)
        os.replace(temp, destination)
    finally:
        if temp.exists():
            temp.unlink()
    return destination
